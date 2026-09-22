"""Testes da planilha configurável: modelo, colunas, importação e exportação."""

import csv
import io

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from ava.models import (
    ColunaPlanilha,
    Curso,
    ImportacaoPlanilha,
    ModeloPlanilha,
    RegistroPlanilha,
    ValorRegistroPlanilha,
)
from ava.services import PlanilhaExportService, PlanilhaImportService
from ava.services.planilha_service import ler_planilha, sugerir_mapeamento
from core.models import AuditLog, Cliente
from curriculum.models import Escola


User = get_user_model()


def montar_xlsx(cabecalhos, linhas) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(list(cabecalhos))
    for linha in linhas:
        sheet.append(list(linha))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def montar_csv(cabecalhos, linhas) -> bytes:
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=";")
    escritor.writerow(cabecalhos)
    escritor.writerows(linhas)
    return buffer.getvalue().encode("utf-8")


@override_settings(
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"]
)
class PlanilhaConfiguravelBaseTests(TestCase):
    def setUp(self):
        self.cliente = Cliente.objects.create(
            nome="Município Planilha", slug="municipio-planilha"
        )
        self.outro_cliente = Cliente.objects.create(
            nome="Outro Município", slug="outro-municipio-planilha"
        )
        self.escola = Escola.objects.create(cliente=self.cliente, nome="Escola Central")
        self.escola_outro = Escola.objects.create(
            cliente=self.outro_cliente, nome="Escola Distante"
        )

        self.admin = User.objects.create_user(
            email="admin-planilha@example.com",
            nome="Admin Planilha",
            password="123456",
            cliente=self.cliente,
            role=User.Role.ADMIN_CLIENTE,
        )
        self.admin_outro = User.objects.create_user(
            email="admin-outro-planilha@example.com",
            nome="Admin Outro",
            password="123456",
            cliente=self.outro_cliente,
            role=User.Role.ADMIN_CLIENTE,
        )
        self.professor = User.objects.create_user(
            email="prof-planilha@example.com",
            nome="Carla Nunes",
            password="123456",
            cliente=self.cliente,
            escola=self.escola,
            role=User.Role.PROFESSOR,
        )
        self.curso = Curso.objects.create(
            cliente=self.cliente,
            titulo="Curso Planilha",
            slug="curso-planilha",
            status=Curso.Status.PUBLICADO,
            autor_principal=self.admin,
        )

    def criar_modelo(self, **kwargs) -> ModeloPlanilha:
        dados = {
            "cliente": self.cliente,
            "nome": "Frequência mensal",
            "slug": "frequencia-mensal",
            "created_by": self.admin,
        }
        dados.update(kwargs)
        return ModeloPlanilha.objects.create(**dados)

    def criar_colunas(self, modelo) -> dict[str, ColunaPlanilha]:
        colunas = {
            "turma": ColunaPlanilha.objects.create(
                cliente=self.cliente,
                modelo=modelo,
                chave="turma",
                titulo="Turma",
                tipo=ColunaPlanilha.Tipo.TEXTO,
                ordem=1,
                obrigatorio=True,
            ),
            "data": ColunaPlanilha.objects.create(
                cliente=self.cliente,
                modelo=modelo,
                chave="data",
                titulo="Data",
                tipo=ColunaPlanilha.Tipo.DATA,
                ordem=2,
            ),
            "presentes": ColunaPlanilha.objects.create(
                cliente=self.cliente,
                modelo=modelo,
                chave="presentes",
                titulo="Presentes",
                tipo=ColunaPlanilha.Tipo.INTEIRO,
                ordem=3,
                validacoes={"min": 0, "max": 40},
            ),
            "turno": ColunaPlanilha.objects.create(
                cliente=self.cliente,
                modelo=modelo,
                chave="turno",
                titulo="Turno",
                tipo=ColunaPlanilha.Tipo.SELECAO,
                ordem=4,
                opcoes=["Manhã", "Tarde"],
            ),
            "professor": ColunaPlanilha.objects.create(
                cliente=self.cliente,
                modelo=modelo,
                chave="professor",
                titulo="Professor",
                tipo=ColunaPlanilha.Tipo.PROFESSOR,
                ordem=5,
            ),
        }
        return colunas

    def criar_importacao(self, modelo, conteudo, nome="planilha.xlsx", modo=None):
        importacao = ImportacaoPlanilha(
            cliente=self.cliente,
            modelo=modelo,
            arquivo=SimpleUploadedFile(nome, conteudo),
            modo=modo or ImportacaoPlanilha.Modo.ADICIONAR,
            created_by=self.admin,
        )
        importacao.save()
        return importacao


class ModeloPlanilhaTests(PlanilhaConfiguravelBaseTests):
    def test_cria_modelo_com_colunas_configuraveis(self):
        modelo = self.criar_modelo()
        colunas = self.criar_colunas(modelo)
        self.assertEqual(modelo.colunas_ativas().count(), 5)
        self.assertEqual(colunas["turno"].opcoes_normalizadas(), ["Manhã", "Tarde"])

    def test_coluna_de_selecao_sem_opcoes_e_invalida(self):
        modelo = self.criar_modelo()
        coluna = ColunaPlanilha(
            cliente=self.cliente,
            modelo=modelo,
            chave="sem-opcoes",
            titulo="Sem opções",
            tipo=ColunaPlanilha.Tipo.SELECAO,
        )
        with self.assertRaises(Exception):
            coluna.full_clean()

    def test_admin_cria_modelo_pela_interface(self):
        self.client.force_login(self.admin)
        resposta = self.client.post(
            reverse("ava:planilha_modelo_nova"),
            {
                "municipio": self.cliente.id,
                "nome": "Acompanhamento semanal",
                "slug": "",
                "descricao": "Registro semanal por turma",
                "versao": "1",
                "ativo": "on",
                "permite_professor": "on",
                "colunas-TOTAL_FORMS": "2",
                "colunas-INITIAL_FORMS": "0",
                "colunas-MIN_NUM_FORMS": "0",
                "colunas-MAX_NUM_FORMS": "1000",
                "colunas-0-titulo": "Turma",
                "colunas-0-chave": "",
                "colunas-0-tipo": ColunaPlanilha.Tipo.TEXTO,
                "colunas-0-ordem": "1",
                "colunas-0-obrigatorio": "on",
                "colunas-0-visivel": "on",
                "colunas-0-valor_padrao": "",
                "colunas-0-ajuda": "",
                "colunas-0-validacoes": "",
                "colunas-0-opcoes_texto": "",
                "colunas-1-titulo": "Turno",
                "colunas-1-chave": "",
                "colunas-1-tipo": ColunaPlanilha.Tipo.SELECAO,
                "colunas-1-ordem": "2",
                "colunas-1-visivel": "on",
                "colunas-1-valor_padrao": "",
                "colunas-1-ajuda": "",
                "colunas-1-validacoes": "",
                "colunas-1-opcoes_texto": "Manhã\nTarde",
            },
        )
        self.assertEqual(resposta.status_code, 302)
        modelo = ModeloPlanilha.objects.get(nome="Acompanhamento semanal")
        self.assertEqual(modelo.slug, "acompanhamento-semanal")
        chaves = set(modelo.colunas_ativas().values_list("chave", flat=True))
        self.assertEqual(chaves, {"turma", "turno"})

    def test_professor_nao_acessa_area_de_configuracao(self):
        self.client.force_login(self.professor)
        resposta = self.client.get(reverse("ava:planilha_modelo_lista"))
        self.assertEqual(resposta.status_code, 403)

    def test_admin_de_outro_municipio_nao_ve_o_modelo(self):
        modelo = self.criar_modelo()
        self.client.force_login(self.admin_outro)
        resposta = self.client.get(
            reverse("ava:planilha_modelo_editar", args=[modelo.id])
        )
        self.assertEqual(resposta.status_code, 404)


class PlanilhaLeituraTests(PlanilhaConfiguravelBaseTests):
    def test_le_cabecalhos_e_linhas_de_xlsx(self):
        conteudo = montar_xlsx(
            ["Turma", "Data", "Presentes"],
            [["3A", "2026-03-01", 20], ["3B", "2026-03-02", 18]],
        )
        leitura = ler_planilha(io.BytesIO(conteudo), "arquivo.xlsx")
        self.assertEqual(leitura.cabecalhos, ["Turma", "Data", "Presentes"])
        self.assertEqual(len(leitura.linhas), 2)
        self.assertEqual(leitura.linhas[0].valores["Turma"], "3A")

    def test_le_csv_com_ponto_e_virgula(self):
        conteudo = montar_csv(["Turma", "Presentes"], [["4A", 25]])
        leitura = ler_planilha(io.BytesIO(conteudo), "arquivo.csv")
        self.assertEqual(leitura.cabecalhos, ["Turma", "Presentes"])
        self.assertEqual(leitura.linhas[0].valores["Presentes"], "25")

    def test_sugere_mapeamento_ignorando_acentos_e_caixa(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        mapeamento = sugerir_mapeamento(modelo, ["TURMA", "data", "Presentes"])
        self.assertEqual(
            mapeamento, {"TURMA": "turma", "data": "data", "Presentes": "presentes"}
        )


class PlanilhaImportacaoTests(PlanilhaConfiguravelBaseTests):
    def test_importacao_valida_grava_registros_consultaveis(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        conteudo = montar_xlsx(
            ["Turma", "Data", "Presentes", "Turno", "Professor"],
            [
                ["3A", "01/03/2026", 20, "Manhã", "prof-planilha@example.com"],
                ["3B", "02/03/2026", 18, "Tarde", "Carla Nunes"],
            ],
        )
        importacao = self.criar_importacao(modelo, conteudo)
        servico = PlanilhaImportService(importacao, self.admin)
        leitura = servico.ler()
        importacao.mapeamento = sugerir_mapeamento(modelo, leitura.cabecalhos)
        importacao.save()

        servico = PlanilhaImportService(importacao, self.admin)
        validacao = servico.validar()
        self.assertTrue(validacao.valida, validacao.erros)

        servico.gravar(validacao)
        importacao.refresh_from_db()
        self.assertEqual(importacao.status, ImportacaoPlanilha.Status.CONCLUIDA)
        self.assertEqual(importacao.total_criados, 2)
        self.assertEqual(RegistroPlanilha.objects.filter(modelo=modelo).count(), 2)

        registro = (
            RegistroPlanilha.objects.filter(modelo=modelo)
            .order_by("linha_origem")
            .first()
        )
        valores = registro.valores_por_chave()
        self.assertEqual(valores["presentes"].valor_numero, 20)
        self.assertEqual(valores["data"].valor_data.isoformat(), "2026-03-01")
        self.assertEqual(valores["professor"].valor_referencia_id, self.professor.pk)
        self.assertEqual(valores["professor"].valor_texto, "Carla Nunes")

    def test_importacao_com_erro_nao_grava_nada(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        conteudo = montar_xlsx(
            ["Turma", "Data", "Presentes", "Turno"],
            [
                ["3A", "01/03/2026", 20, "Manhã"],
                ["", "data-invalida", 99, "Noite"],
            ],
        )
        importacao = self.criar_importacao(modelo, conteudo)
        servico = PlanilhaImportService(importacao, self.admin)
        leitura = servico.ler()
        importacao.mapeamento = sugerir_mapeamento(modelo, leitura.cabecalhos)
        importacao.save()

        servico = PlanilhaImportService(importacao, self.admin)
        validacao = servico.validar()
        self.assertFalse(validacao.valida)

        colunas_com_erro = {erro.coluna for erro in validacao.erros}
        self.assertIn("Turma", colunas_com_erro)
        self.assertIn("Data", colunas_com_erro)
        self.assertIn("Presentes", colunas_com_erro)
        self.assertIn("Turno", colunas_com_erro)
        self.assertTrue(all(erro.linha == 3 for erro in validacao.erros))

        with self.assertRaises(ValueError):
            servico.gravar(validacao)
        self.assertEqual(RegistroPlanilha.objects.filter(modelo=modelo).count(), 0)

    def test_coluna_obrigatoria_nao_mapeada_bloqueia_a_importacao(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        conteudo = montar_xlsx(["Presentes"], [[10]])
        importacao = self.criar_importacao(modelo, conteudo)
        importacao.mapeamento = {"Presentes": "presentes"}
        importacao.save()

        validacao = PlanilhaImportService(importacao, self.admin).validar()
        self.assertFalse(validacao.valida)
        self.assertTrue(
            any("não foi mapeada" in erro.mensagem for erro in validacao.erros)
        )

    def test_modo_atualizar_reaproveita_registro_pela_chave(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        conteudo = montar_xlsx(["Turma", "Presentes"], [["3A", 20]])

        primeira = self.criar_importacao(modelo, conteudo)
        primeira.mapeamento = {"Turma": "turma", "Presentes": "presentes"}
        primeira.coluna_chave = "turma"
        primeira.save()
        servico = PlanilhaImportService(primeira, self.admin)
        servico.gravar(servico.validar())
        self.assertEqual(RegistroPlanilha.objects.filter(modelo=modelo).count(), 1)

        atualizado = montar_xlsx(["Turma", "Presentes"], [["3A", 27]])
        segunda = self.criar_importacao(
            modelo, atualizado, modo=ImportacaoPlanilha.Modo.ATUALIZAR
        )
        segunda.mapeamento = {"Turma": "turma", "Presentes": "presentes"}
        segunda.coluna_chave = "turma"
        segunda.save()
        servico = PlanilhaImportService(segunda, self.admin)
        servico.gravar(servico.validar())

        segunda.refresh_from_db()
        self.assertEqual(segunda.total_atualizados, 1)
        self.assertEqual(segunda.total_criados, 0)
        self.assertEqual(RegistroPlanilha.objects.filter(modelo=modelo).count(), 1)
        registro = RegistroPlanilha.objects.get(modelo=modelo)
        self.assertEqual(registro.valores_por_chave()["presentes"].valor_numero, 27)

    def test_modo_atualizar_exige_coluna_chave(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        conteudo = montar_xlsx(["Turma"], [["3A"]])
        importacao = self.criar_importacao(
            modelo, conteudo, modo=ImportacaoPlanilha.Modo.ATUALIZAR
        )
        importacao.mapeamento = {"Turma": "turma"}
        importacao.save()

        validacao = PlanilhaImportService(importacao, self.admin).validar()
        self.assertFalse(validacao.valida)

    def test_relacional_inexistente_gera_erro_de_linha(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        conteudo = montar_xlsx(
            ["Turma", "Professor"], [["3A", "fantasma@example.com"]]
        )
        importacao = self.criar_importacao(modelo, conteudo)
        importacao.mapeamento = {"Turma": "turma", "Professor": "professor"}
        importacao.save()

        validacao = PlanilhaImportService(importacao, self.admin).validar()
        self.assertFalse(validacao.valida)
        self.assertTrue(
            any("não encontrado" in erro.mensagem for erro in validacao.erros)
        )

    def test_importacao_registra_auditoria(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        conteudo = montar_xlsx(["Turma"], [["3A"]])
        importacao = self.criar_importacao(modelo, conteudo)
        importacao.mapeamento = {"Turma": "turma"}
        importacao.save()
        servico = PlanilhaImportService(importacao, self.admin)
        servico.gravar(servico.validar())

        self.assertTrue(
            AuditLog.objects.filter(
                cliente=self.cliente,
                entidade="ava.ImportacaoPlanilha",
                acao="importacao_concluida",
            ).exists()
        )


class PlanilhaExportacaoTests(PlanilhaConfiguravelBaseTests):
    def _modelo_com_um_registro(self):
        modelo = self.criar_modelo()
        colunas = self.criar_colunas(modelo)
        registro = RegistroPlanilha.objects.create(
            cliente=self.cliente,
            modelo=modelo,
            escola=self.escola,
            chave_externa="3A",
            created_by=self.admin,
        )
        ValorRegistroPlanilha.objects.create(
            cliente=self.cliente,
            registro=registro,
            coluna=colunas["turma"],
            valor_texto="3A",
        )
        ValorRegistroPlanilha.objects.create(
            cliente=self.cliente,
            registro=registro,
            coluna=colunas["presentes"],
            valor_numero=21,
        )
        return modelo, registro

    def test_exportacao_csv_traz_os_valores_gravados(self):
        modelo, registro = self._modelo_com_um_registro()
        payload = PlanilhaExportService.para_csv(modelo, [registro])
        conteudo = payload.decode("utf-8-sig")
        self.assertIn("Turma", conteudo)
        self.assertIn("3A", conteudo)
        self.assertIn("21", conteudo)

    def test_exportacao_xlsx_pela_interface(self):
        modelo, _ = self._modelo_com_um_registro()
        self.client.force_login(self.admin)
        resposta = self.client.get(
            reverse("ava:planilha_exportar", args=[modelo.id, "xlsx"])
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("spreadsheetml", resposta["Content-Type"])

    def test_planilha_modelo_em_branco_traz_apenas_cabecalhos(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        payload = PlanilhaExportService.modelo_em_branco_xlsx(modelo)
        workbook = openpyxl.load_workbook(io.BytesIO(payload))
        sheet = workbook.active
        cabecalhos = [celula.value for celula in sheet[1]]
        self.assertEqual(
            cabecalhos, ["Turma", "Data", "Presentes", "Turno", "Professor"]
        )
        self.assertEqual(sheet.max_row, 1)


class PlanilhaPermissaoTests(PlanilhaConfiguravelBaseTests):
    def test_professor_ve_apenas_modelos_ativos_liberados(self):
        liberado = self.criar_modelo(
            nome="Liberado", slug="liberado", permite_professor=True
        )
        self.criar_modelo(
            nome="Restrito", slug="restrito", permite_professor=False
        )
        self.client.force_login(self.professor)
        resposta = self.client.get(
            reverse("ava:planilha_registros", args=[liberado.id])
        )
        self.assertEqual(resposta.status_code, 200)

    def test_professor_nao_acessa_modelo_restrito(self):
        restrito = self.criar_modelo(
            nome="Restrito", slug="restrito", permite_professor=False
        )
        self.client.force_login(self.professor)
        resposta = self.client.get(
            reverse("ava:planilha_registros", args=[restrito.id])
        )
        self.assertEqual(resposta.status_code, 404)

    def test_modelo_de_outra_escola_nao_aparece_para_o_professor(self):
        outra_escola = Escola.objects.create(cliente=self.cliente, nome="Escola Norte")
        modelo = self.criar_modelo(
            nome="Só da escola Norte", slug="escola-norte", escola=outra_escola
        )
        self.client.force_login(self.professor)
        resposta = self.client.get(
            reverse("ava:planilha_registros", args=[modelo.id])
        )
        self.assertEqual(resposta.status_code, 404)

    def test_isolamento_entre_municipios_na_importacao(self):
        modelo = self.criar_modelo()
        self.client.force_login(self.admin_outro)
        resposta = self.client.get(
            reverse("ava:planilha_importar", args=[modelo.id])
        )
        self.assertEqual(resposta.status_code, 404)


class PlanilhaFluxoInterfaceTests(PlanilhaConfiguravelBaseTests):
    def test_fluxo_completo_de_importacao_pela_interface(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        conteudo = montar_xlsx(
            ["Turma", "Data", "Presentes", "Turno"],
            [["3A", "01/03/2026", 20, "Manhã"]],
        )
        self.client.force_login(self.admin)

        upload = self.client.post(
            reverse("ava:planilha_importar", args=[modelo.id]),
            {
                "arquivo": SimpleUploadedFile("dados.xlsx", conteudo),
                "modo": ImportacaoPlanilha.Modo.ADICIONAR,
            },
        )
        self.assertEqual(upload.status_code, 302)
        importacao = ImportacaoPlanilha.objects.get(modelo=modelo)
        self.assertEqual(
            importacao.cabecalhos, ["Turma", "Data", "Presentes", "Turno"]
        )

        url_mapear = reverse("ava:planilha_importar_mapear", args=[importacao.id])
        pagina = self.client.get(url_mapear)
        self.assertEqual(pagina.status_code, 200)
        self.assertIn("Pré-visualização", pagina.content.decode())

        dados_mapeamento = {
            "acao": "validar",
            "col_0": "turma",
            "col_1": "data",
            "col_2": "presentes",
            "col_3": "turno",
            "coluna_chave": "",
        }
        validacao = self.client.post(url_mapear, dados_mapeamento)
        self.assertEqual(validacao.status_code, 200)
        self.assertEqual(RegistroPlanilha.objects.filter(modelo=modelo).count(), 0)

        dados_mapeamento["acao"] = "confirmar"
        confirmacao = self.client.post(url_mapear, dados_mapeamento)
        self.assertEqual(confirmacao.status_code, 302)
        self.assertEqual(RegistroPlanilha.objects.filter(modelo=modelo).count(), 1)

    def test_planilha_invalida_nao_e_gravada_silenciosamente(self):
        modelo = self.criar_modelo()
        self.criar_colunas(modelo)
        conteudo = montar_xlsx(["Turma", "Presentes"], [["3A", "muitos"]])
        self.client.force_login(self.admin)

        self.client.post(
            reverse("ava:planilha_importar", args=[modelo.id]),
            {
                "arquivo": SimpleUploadedFile("dados.xlsx", conteudo),
                "modo": ImportacaoPlanilha.Modo.ADICIONAR,
            },
        )
        importacao = ImportacaoPlanilha.objects.get(modelo=modelo)
        resposta = self.client.post(
            reverse("ava:planilha_importar_mapear", args=[importacao.id]),
            {
                "acao": "confirmar",
                "col_0": "turma",
                "col_1": "presentes",
                "coluna_chave": "",
            },
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(RegistroPlanilha.objects.filter(modelo=modelo).count(), 0)
        importacao.refresh_from_db()
        self.assertEqual(importacao.status, ImportacaoPlanilha.Status.ERRO)
        self.assertTrue(importacao.erros)

    def test_busca_e_ordenacao_nos_registros(self):
        modelo = self.criar_modelo()
        colunas = self.criar_colunas(modelo)
        for turma, presentes in (("3A", 10), ("3B", 30)):
            registro = RegistroPlanilha.objects.create(
                cliente=self.cliente,
                modelo=modelo,
                created_by=self.admin,
                chave_externa=turma,
            )
            ValorRegistroPlanilha.objects.create(
                cliente=self.cliente,
                registro=registro,
                coluna=colunas["turma"],
                valor_texto=turma,
            )
            ValorRegistroPlanilha.objects.create(
                cliente=self.cliente,
                registro=registro,
                coluna=colunas["presentes"],
                valor_numero=presentes,
            )

        self.client.force_login(self.admin)
        busca = self.client.get(
            reverse("ava:planilha_registros", args=[modelo.id]), {"q": "3B"}
        )
        conteudo = busca.content.decode()
        self.assertIn("3B", conteudo)

        ordenado = self.client.get(
            reverse("ava:planilha_registros", args=[modelo.id]),
            {"ordenar": "presentes", "direcao": "desc"},
        )
        self.assertEqual(ordenado.status_code, 200)
