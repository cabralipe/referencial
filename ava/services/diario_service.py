"""Regras de negócio do Diário de Bordo."""

from __future__ import annotations

import io

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.utils import timezone

from ava.models.diario import (
    DiarioBordo,
    DiarioBordoMidia,
    DiarioBordoParticipante,
    MAX_DIARIO_MIDIA_SIZE,
)
from ava.services.audit_service import AVAAuditService, diff_de_campos
from core.models import Usuario


CAMPOS_AUDITADOS = [
    "origem",
    "data_aula",
    "titulo",
    "status",
    "escola_id",
    "curso_id",
    "modulo_id",
    "aula_id",
    "professor_id",
    "participantes_presentes",
]

#: Papéis que podem revisar/bloquear registros de terceiros.
ROLES_REVISAO = {
    Usuario.Role.ADMIN_CLIENTE,
    Usuario.Role.SUPER_ADMIN,
    Usuario.Role.COORDENADOR_PEDAGOGICO,
    Usuario.Role.DIRETOR,
}


class DiarioBordoService:
    """Transições de status, frequência e mídias do Diário de Bordo."""

    # -- criação/edição -------------------------------------------------
    @staticmethod
    @transaction.atomic
    def salvar(diario: DiarioBordo, usuario, *, criando: bool) -> DiarioBordo:
        diario.updated_by = usuario
        if criando:
            diario.created_by = usuario
        diario.full_clean(exclude=None)
        diario.save()
        diario.sincronizar_frequencia()
        if diario.frequencia_nominal:
            DiarioBordo.raw_objects.filter(pk=diario.pk).update(
                participantes_presentes=diario.participantes_presentes
            )
        AVAAuditService.registrar_diario(
            diario,
            usuario,
            acao="criado" if criando else "atualizado",
            diff=diff_de_campos(diario, CAMPOS_AUDITADOS),
        )
        return diario

    # -- fluxo de status ------------------------------------------------
    @staticmethod
    def pode_editar(diario: DiarioBordo, usuario) -> bool:
        if usuario.role in ROLES_REVISAO:
            return diario.status != DiarioBordo.Status.BLOQUEADO
        return diario.editavel and diario.professor_id == usuario.id

    @staticmethod
    @transaction.atomic
    def enviar(diario: DiarioBordo, usuario) -> DiarioBordo:
        if diario.status != DiarioBordo.Status.RASCUNHO:
            raise ValidationError("Somente rascunhos podem ser enviados.")
        if diario.professor_id != usuario.id and usuario.role not in ROLES_REVISAO:
            raise PermissionDenied("Apenas o professor responsável pode enviar o registro.")
        diario.sincronizar_frequencia()
        if diario.frequencia_nominal and diario.participantes_presentes == 0:
            raise ValidationError(
                "Registre ao menos um participante presente antes de enviar."
            )
        if not diario.frequencia_nominal and diario.participantes_presentes == 0:
            raise ValidationError(
                "Informe a quantidade de participantes presentes antes de enviar."
            )
        diario.status = DiarioBordo.Status.ENVIADO
        diario.enviado_em = timezone.now()
        diario.updated_by = usuario
        diario.save(
            update_fields=[
                "status",
                "enviado_em",
                "updated_by",
                "participantes_presentes",
                "updated_at",
            ]
        )
        AVAAuditService.registrar_diario(diario, usuario, acao="enviado")
        return diario

    @staticmethod
    @transaction.atomic
    def revisar(diario: DiarioBordo, usuario, *, parecer: str = "", bloquear: bool = False):
        if usuario.role not in ROLES_REVISAO:
            raise PermissionDenied("Você não pode revisar registros do diário de bordo.")
        if diario.status == DiarioBordo.Status.RASCUNHO:
            raise ValidationError("O registro ainda é um rascunho e não pode ser revisado.")
        diario.status = (
            DiarioBordo.Status.BLOQUEADO if bloquear else DiarioBordo.Status.REVISADO
        )
        diario.revisado_em = timezone.now()
        diario.revisado_por = usuario
        diario.parecer_revisao = parecer or ""
        diario.updated_by = usuario
        diario.save(
            update_fields=[
                "status",
                "revisado_em",
                "revisado_por",
                "parecer_revisao",
                "updated_by",
                "updated_at",
            ]
        )
        AVAAuditService.registrar_diario(
            diario,
            usuario,
            acao="bloqueado" if bloquear else "revisado",
            diff={"parecer": parecer or ""},
        )
        return diario

    @staticmethod
    @transaction.atomic
    def reabrir(diario: DiarioBordo, usuario) -> DiarioBordo:
        if usuario.role not in ROLES_REVISAO:
            raise PermissionDenied("Você não pode reabrir registros do diário de bordo.")
        diario.status = DiarioBordo.Status.RASCUNHO
        diario.enviado_em = None
        diario.updated_by = usuario
        diario.save(update_fields=["status", "enviado_em", "updated_by", "updated_at"])
        AVAAuditService.registrar_diario(diario, usuario, acao="reaberto")
        return diario

    # -- frequência -----------------------------------------------------
    @staticmethod
    @transaction.atomic
    def registrar_presenca(
        diario: DiarioBordo,
        usuario,
        *,
        nome: str,
        identificacao: str = "",
        presente: bool = True,
        justificativa: str = "",
        usuario_vinculado=None,
    ) -> DiarioBordoParticipante:
        if not DiarioBordoService.pode_editar(diario, usuario):
            raise PermissionDenied("Este registro não pode mais ser alterado.")
        participante = DiarioBordoParticipante(
            cliente_id=diario.cliente_id,
            diario=diario,
            nome=nome,
            identificacao=identificacao,
            presente=presente,
            justificativa=justificativa,
            usuario=usuario_vinculado,
        )
        participante.full_clean()
        participante.save()
        DiarioBordoService._resincronizar(diario, usuario)
        return participante

    @staticmethod
    @transaction.atomic
    def remover_presenca(diario: DiarioBordo, usuario, participante_id: int) -> None:
        if not DiarioBordoService.pode_editar(diario, usuario):
            raise PermissionDenied("Este registro não pode mais ser alterado.")
        participante = DiarioBordoParticipante.raw_objects.filter(
            is_deleted=False, diario=diario, pk=participante_id
        ).first()
        if participante is None:
            return
        participante.delete()
        DiarioBordoService._resincronizar(diario, usuario)

    @staticmethod
    def _resincronizar(diario: DiarioBordo, usuario) -> None:
        diario.sincronizar_frequencia()
        if diario.frequencia_nominal:
            DiarioBordo.raw_objects.filter(pk=diario.pk).update(
                participantes_presentes=diario.participantes_presentes
            )
        AVAAuditService.registrar_diario(
            diario,
            usuario,
            acao="frequencia_atualizada",
            diff={"presentes": diario.participantes_presentes},
        )

    # -- mídias ---------------------------------------------------------
    @staticmethod
    @transaction.atomic
    def adicionar_midia(
        diario: DiarioBordo,
        usuario,
        arquivo,
        *,
        tipo: str = DiarioBordoMidia.Tipo.FOTO,
        legenda: str = "",
        consentimento: bool = False,
    ) -> DiarioBordoMidia:
        if not DiarioBordoService.pode_editar(diario, usuario):
            raise PermissionDenied("Este registro não pode mais receber anexos.")
        if getattr(arquivo, "size", 0) > MAX_DIARIO_MIDIA_SIZE:
            raise ValidationError({"arquivo": "O arquivo deve ter no máximo 10 MB."})

        midia = DiarioBordoMidia(
            cliente_id=diario.cliente_id,
            diario=diario,
            tipo=tipo,
            arquivo=arquivo,
            legenda=legenda,
            consentimento_registrado=consentimento,
            created_by=usuario,
            ordem=(
                DiarioBordoMidia.raw_objects.filter(
                    is_deleted=False, diario=diario
                ).count()
            ),
        )
        midia.full_clean(exclude=["nome_original", "content_type", "tamanho_bytes"])
        midia.save()
        AVAAuditService.registrar_diario(
            diario,
            usuario,
            acao="midia_adicionada",
            diff={"midia": midia.pk, "tipo": tipo, "nome": midia.nome_original},
        )
        return midia

    @staticmethod
    @transaction.atomic
    def remover_midia(diario: DiarioBordo, usuario, midia_id: int) -> None:
        if not DiarioBordoService.pode_editar(diario, usuario):
            raise PermissionDenied("Este registro não pode mais ser alterado.")
        midia = DiarioBordoMidia.raw_objects.filter(
            is_deleted=False, diario=diario, pk=midia_id
        ).first()
        if midia is None:
            return
        midia.delete()
        AVAAuditService.registrar_diario(
            diario, usuario, acao="midia_removida", diff={"midia": midia_id}
        )


class DiarioBordoExportService:
    """Exportação dos registros do diário, identificando a modalidade."""

    CABECALHOS = [
        "Data",
        "Origem",
        "Título",
        "Escola",
        "Professor",
        "Curso",
        "Módulo",
        "Aula",
        "Tipo de atividade",
        "Local",
        "Instituição",
        "Grupo participante",
        "Previstos",
        "Presentes",
        "Status",
        "Conteúdo trabalhado",
        "Metodologia",
        "Observações",
        "Fotos",
        "Criado em",
    ]

    @classmethod
    def montar_linhas(cls, diarios) -> list[list[str]]:
        linhas = []
        for diario in diarios:
            linhas.append(
                [
                    diario.data_aula.strftime("%d/%m/%Y") if diario.data_aula else "",
                    diario.get_origem_display(),
                    diario.titulo,
                    diario.escola.nome if diario.escola_id else "",
                    (diario.professor.nome or diario.professor.email)
                    if diario.professor_id
                    else "",
                    diario.curso.titulo if diario.curso_id else "",
                    diario.modulo.titulo if diario.modulo_id else "",
                    diario.aula.titulo if diario.aula_id else "",
                    diario.get_tipo_atividade_display() if diario.tipo_atividade else "",
                    diario.local_realizacao,
                    diario.instituicao_parceira,
                    diario.grupo_participante,
                    diario.participantes_previstos
                    if diario.participantes_previstos is not None
                    else "",
                    diario.participantes_presentes,
                    diario.get_status_display(),
                    diario.conteudo_trabalhado,
                    diario.metodologia,
                    diario.observacoes,
                    diario.midias.filter(is_deleted=False).count(),
                    diario.created_at.strftime("%d/%m/%Y %H:%M"),
                ]
            )
        return linhas

    @classmethod
    def para_xlsx(cls, diarios) -> bytes:
        import openpyxl
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Diário de bordo"
        sheet.append(cls.CABECALHOS)
        for celula in sheet[1]:
            celula.font = Font(bold=True)
        for linha in cls.montar_linhas(diarios):
            sheet.append(linha)
        for indice, cabecalho in enumerate(cls.CABECALHOS, start=1):
            sheet.column_dimensions[get_column_letter(indice)].width = min(
                max(len(cabecalho) + 4, 14), 45
            )
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @classmethod
    def para_csv(cls, diarios) -> bytes:
        import csv

        buffer = io.StringIO()
        escritor = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
        escritor.writerow(cls.CABECALHOS)
        escritor.writerows(cls.montar_linhas(diarios))
        return buffer.getvalue().encode("utf-8-sig")
