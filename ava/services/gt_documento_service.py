"""Geração dos documentos DOCX dos GTs.

Três documentos, todos em DOCX real (python-docx, a mesma dependência já usada
pelo plugin de exportação):

1. documento do objetivo do GT — objetivo oficial, proposta do GT, diagnóstico,
   metas e estratégias;
2. ata de uma reunião;
3. consolidação institucional com todos os objetivos na ordem oficial.

Cada geração fica registrada em :class:`exports.models.ExportJob` (o mecanismo
de exportação que já existe no projeto) e o arquivo vai para o storage privado —
o download passa por view autorizada, nunca por URL pública.
"""

from __future__ import annotations

import re
import unicodedata
from io import BytesIO

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.base import ContentFile
from django.db.models import Q
from django.utils import timezone

from ava.models.gt_sala import (
    AtaReuniaoGT,
    CicloReferencial,
    EstrategiaGT,
    MetaGT,
    PresencaReuniaoGT,
    ReuniaoGT,
    SalaGT,
    TrabalhoGT,
)
from ava.services.audit_service import AVAAuditService
from ava.services.gt_access import acesso_a_sala, tem_visao_global
from ava.services.gt_reuniao_service import AtaGTService
from ava.storage import private_ava_storage
from core.models import ClienteTema
from exports.models import ExportJob
from exports.services.docx import html_para_paragrafos


try:  # pragma: no cover - dependência externa já presente em requirements
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Cm, Pt
except ImportError:  # pragma: no cover
    Document = None  # type: ignore


PASTA_DOCUMENTOS = "ava/gt-documentos"

#: Controles que o XML do DOCX não aceita.
_CONTROLES = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

TEXTO_PENDENTE = "Pendente de consolidação"


def _limpar(texto) -> str:
    """Normaliza o texto para o XML do DOCX preservando a acentuação."""

    if texto is None:
        return ""
    return _CONTROLES.sub("", str(texto)).replace("\r\n", "\n").strip()


def _slug(valor: str, tamanho: int = 40) -> str:
    base = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    base = re.sub(r"[^a-zA-Z0-9]+", "-", base).strip("-").lower()
    return (base or "documento")[:tamanho]


def _campo(paragrafo, instrucao: str) -> None:
    """Insere um campo do Word (sumário, número de página) no parágrafo."""

    run = paragrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instrucao
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    fim = OxmlElement("w:fldChar")
    fim.set(qn("w:fldCharType"), "end")
    for elemento in (inicio, instr, separador, fim):
        run._r.append(elemento)


def _definir_idioma(document) -> None:
    estilo = document.styles["Normal"]
    rpr = estilo.element.get_or_add_rPr()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "pt-BR")


def _tema_do_cliente(cliente_id: int) -> dict:
    tema = (
        ClienteTema.raw_objects.filter(cliente_id=cliente_id, is_deleted=False)
        .values("cabecalho_html", "rodape_html")
        .first()
        or {}
    )
    return {
        "cabecalho": " ".join(
            texto for _, texto in html_para_paragrafos(tema.get("cabecalho_html"))
        ),
        "rodape": " ".join(texto for _, texto in html_para_paragrafos(tema.get("rodape_html"))),
    }


def _novo_documento(cliente_id: int, titulo: str) -> "Document":
    if Document is None:  # pragma: no cover
        raise RuntimeError("python-docx não está disponível no ambiente")

    document = Document()
    _definir_idioma(document)
    document.core_properties.title = _limpar(titulo)
    document.core_properties.language = "pt-BR"

    for secao in document.sections:
        secao.left_margin = Cm(2.5)
        secao.right_margin = Cm(2.5)
        secao.top_margin = Cm(2)
        secao.bottom_margin = Cm(2)

    tema = _tema_do_cliente(cliente_id)
    if tema["cabecalho"]:
        cabecalho = document.sections[0].header.paragraphs[0]
        cabecalho.text = _limpar(tema["cabecalho"])[:300]
        cabecalho.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cabecalho.runs:
            run.font.size = Pt(9)

    rodape = document.sections[0].footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefixo = _limpar(tema["rodape"])[:200]
    if prefixo:
        rodape.add_run(f"{prefixo} — ")
    rodape.add_run("Página ")
    _campo(rodape, "PAGE")
    rodape.add_run(" de ")
    _campo(rodape, "NUMPAGES")
    for run in rodape.runs:
        run.font.size = Pt(8)
    return document


def _paragrafo_texto(document, texto: str, *, vazio: str = "Não informado.") -> None:
    """Escreve um texto multilinha preservando listas simples e campos vazios."""

    conteudo = _limpar(texto)
    if not conteudo:
        paragrafo = document.add_paragraph(vazio)
        for run in paragrafo.runs:
            run.italic = True
        return
    for linha in conteudo.split("\n"):
        linha = linha.strip()
        if not linha:
            continue
        if linha.startswith(("- ", "* ", "• ")):
            paragrafo = document.add_paragraph(linha[2:].strip())
            try:
                paragrafo.style = "List Bullet"
            except (KeyError, AttributeError):  # pragma: no cover - estilo ausente
                pass
        else:
            document.add_paragraph(linha)


def _tabela(document, cabecalhos: list[str], linhas: list[list[str]], *, larguras=None) -> None:
    tabela = document.add_table(rows=1, cols=len(cabecalhos))
    try:
        tabela.style = "Table Grid"
    except KeyError:  # pragma: no cover - estilo ausente
        pass
    tabela.autofit = True
    cabecalho = tabela.rows[0].cells
    for indice, titulo in enumerate(cabecalhos):
        cabecalho[indice].text = titulo
        for paragrafo in cabecalho[indice].paragraphs:
            for run in paragrafo.runs:
                run.bold = True
    for linha in linhas:
        celulas = tabela.add_row().cells
        for indice, valor in enumerate(linha):
            celulas[indice].text = _limpar(valor) or "—"
    if larguras:
        for indice, largura in enumerate(larguras):
            for linha in tabela.rows:
                if indice < len(linha.cells):
                    linha.cells[indice].width = Cm(largura)


def _ficha(document, itens: list[tuple[str, str]]) -> None:
    tabela = document.add_table(rows=0, cols=2)
    try:
        tabela.style = "Table Grid"
    except KeyError:  # pragma: no cover
        pass
    for rotulo, valor in itens:
        celulas = tabela.add_row().cells
        celulas[0].text = rotulo
        celulas[1].text = _limpar(valor) or "—"
        for paragrafo in celulas[0].paragraphs:
            for run in paragrafo.runs:
                run.bold = True
        celulas[0].width = Cm(5)
        celulas[1].width = Cm(11)


def _coordenacao(sala: SalaGT) -> str:
    from ava.services.gt_sala_service import SalaGTService

    nomes = [
        inscricao.usuario.nome or inscricao.usuario.email
        for inscricao in SalaGTService.membros_ativos(sala)
        if inscricao.papel == inscricao.Papel.COORDENACAO
    ]
    return ", ".join(nomes)


def _metas_e_estrategias(trabalho: TrabalhoGT) -> tuple[list[MetaGT], list[EstrategiaGT]]:
    metas = list(
        MetaGT.raw_objects.filter(is_deleted=False, trabalho_id=trabalho.pk).order_by(
            "ordem", "id"
        )
    )
    estrategias = list(
        EstrategiaGT.raw_objects.filter(is_deleted=False, trabalho_id=trabalho.pk)
        .order_by("ordem", "id")
        .prefetch_related("metas")
    )
    return metas, estrategias


def _secao_objetivo(document, sala: SalaGT, trabalho: TrabalhoGT | None, *, somente_consolidavel: bool):
    """Escreve o bloco completo de um objetivo. Devolve ``True`` se houve conteúdo."""

    objetivo = sala.objetivo
    document.add_heading(f"Objetivo {objetivo.codigo} — {_limpar(objetivo.titulo)}", level=1)

    _ficha(
        document,
        [
            ("Ciclo", sala.ciclo.nome),
            ("Grupo de Trabalho", sala.gt.nome),
            ("Situação da sala", sala.get_situacao_display()),
            (
                "Situação do conteúdo",
                trabalho.get_estado_display() if trabalho else "Sem conteúdo iniciado",
            ),
            ("Coordenação do GT", _coordenacao(sala)),
        ],
    )

    document.add_heading("Redação oficial do objetivo", level=2)
    _paragrafo_texto(document, objetivo.texto, vazio="Redação oficial não cadastrada.")

    if trabalho is None or (somente_consolidavel and not trabalho.consolidavel):
        document.add_heading("Situação", level=2)
        paragrafo = document.add_paragraph(
            f"{TEXTO_PENDENTE}: o conteúdo deste GT ainda não foi aprovado."
        )
        for run in paragrafo.runs:
            run.bold = True
        return False

    document.add_heading("Proposta de redação do GT", level=2)
    _paragrafo_texto(
        document,
        trabalho.objetivo_proposto,
        vazio="O GT não propôs nova redação; vale a redação oficial.",
    )

    document.add_heading("Diagnóstico", level=2)
    if trabalho.diagnostico_nao_aplicavel:
        document.add_paragraph(
            "O GT registrou que o diagnóstico não se aplica a este objetivo."
        )
    else:
        for rotulo, campo in (
            ("Caracterização da situação atual", trabalho.diagnostico_situacao),
            ("Evidências e fontes", trabalho.diagnostico_evidencias),
            ("Desafios", trabalho.diagnostico_desafios),
            ("Potencialidades", trabalho.diagnostico_potencialidades),
            ("Público e território", trabalho.diagnostico_publico),
        ):
            document.add_heading(rotulo, level=3)
            _paragrafo_texto(document, campo)

    metas, estrategias = _metas_e_estrategias(trabalho)

    document.add_heading("Metas e resultados esperados", level=2)
    if metas:
        _tabela(
            document,
            ["#", "Meta", "Indicador/evidência", "Prazo", "Observações"],
            [
                [
                    str(indice),
                    meta.descricao,
                    meta.indicador,
                    meta.prazo,
                    meta.observacoes,
                ]
                for indice, meta in enumerate(metas, start=1)
            ],
            larguras=[1, 6, 4.5, 2.5, 3],
        )
    else:
        _paragrafo_texto(document, "", vazio="Nenhuma meta registrada.")

    document.add_heading("Estratégias e ações", level=2)
    if estrategias:
        posicao_das_metas = {meta.pk: indice for indice, meta in enumerate(metas, start=1)}
        linhas = []
        for indice, estrategia in enumerate(estrategias, start=1):
            vinculos = ", ".join(
                f"Meta {posicao_das_metas[meta.pk]}"
                for meta in estrategia.metas.all()
                if meta.pk in posicao_das_metas
            )
            linhas.append(
                [
                    str(indice),
                    estrategia.descricao,
                    estrategia.responsavel,
                    estrategia.prazo,
                    estrategia.recursos,
                    estrategia.indicador,
                    vinculos,
                ]
            )
        _tabela(
            document,
            [
                "#",
                "Estratégia/ação",
                "Responsável",
                "Prazo",
                "Recursos/apoios",
                "Indicador/evidência",
                "Metas",
            ],
            linhas,
            larguras=[0.8, 4.2, 2.4, 1.8, 2.4, 2.4, 2],
        )
    else:
        _paragrafo_texto(document, "", vazio="Nenhuma estratégia registrada.")

    if _limpar(trabalho.observacoes) or _limpar(trabalho.pendencias):
        document.add_heading("Observações e pendências", level=2)
        _paragrafo_texto(document, trabalho.observacoes, vazio="Sem observações.")
        _paragrafo_texto(document, trabalho.pendencias, vazio="Sem pendências registradas.")
    return True


class DocumentoGTService:
    """Monta os DOCX e registra cada geração como ``ExportJob``."""

    # ------------------------------------------------------------------
    # Renderização
    # ------------------------------------------------------------------
    @staticmethod
    def render_objetivo(sala: SalaGT, trabalho: TrabalhoGT | None) -> bytes:
        document = _novo_documento(
            sala.cliente_id,
            f"Objetivo {sala.objetivo.codigo} - {sala.gt.nome}",
        )
        document.add_heading("Documento do objetivo do Referencial", level=0)
        document.add_paragraph(
            f"Município: {sala.cliente.nome} · Ciclo: {_limpar(sala.ciclo.nome)} · "
            f"Emitido em {timezone.localtime():%d/%m/%Y %H:%M}"
        )
        _secao_objetivo(document, sala, trabalho, somente_consolidavel=False)
        if trabalho is not None:
            document.add_paragraph(
                f"Versão interna do conteúdo: v{trabalho.version} · "
                f"Situação: {trabalho.get_estado_display()}"
            )
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def render_ata(reuniao: ReuniaoGT, ata: AtaReuniaoGT) -> bytes:
        sala = reuniao.sala
        document = _novo_documento(
            sala.cliente_id, f"Ata {sala.objetivo.codigo} - {reuniao.data_hora:%Y-%m-%d}"
        )
        titulo = "Ata de reunião" if ata.validada else "Minuta de ata de reunião"
        document.add_heading(titulo, level=0)
        if not ata.validada:
            aviso = document.add_paragraph(
                "Documento preliminar: esta minuta ainda não foi validada pela coordenação do GT."
            )
            for run in aviso.runs:
                run.bold = True

        _ficha(
            document,
            [
                ("Município", sala.cliente.nome),
                ("Ciclo", sala.ciclo.nome),
                ("Objetivo", f"{sala.objetivo.codigo} — {sala.objetivo.titulo}"),
                ("Grupo de Trabalho", sala.gt.nome),
                ("Reunião", reuniao.titulo),
                ("Data e hora", f"{timezone.localtime(reuniao.data_hora):%d/%m/%Y %H:%M}"),
                ("Modalidade", reuniao.get_modalidade_display()),
                ("Local/acesso", reuniao.local or reuniao.link),
                (
                    "Responsável pelo registro",
                    reuniao.responsavel_registro.nome or reuniao.responsavel_registro.email,
                ),
                ("Modelo de ata", f"{ata.modelo_nome} (v{ata.modelo_versao})"),
            ],
        )

        document.add_heading("Participantes", level=1)
        participantes = AtaGTService.participantes_para_ata(reuniao)
        if participantes:
            # A ata registra presença/ausência, nunca a justificativa pessoal.
            _tabela(
                document,
                ["Nome", "Papel no GT", "Presença"],
                [
                    [item["nome"], item["papel"], "Presente" if item["presente"] else "Ausente"]
                    for item in participantes
                ],
                larguras=[7, 5, 4],
            )
        else:
            _paragrafo_texto(document, "", vazio="Frequência não lançada.")

        secoes = (ata.estrutura_json or {}).get("secoes") or []
        mapa = {
            "pauta": ata.pauta,
            "sintese": ata.sintese,
            "decisoes": ata.decisoes,
            "encaminhamentos": ata.encaminhamentos,
            "pendencias": ata.pendencias,
            "assinaturas": ata.assinaturas,
        }
        extras = ata.secoes_extra if isinstance(ata.secoes_extra, dict) else {}
        for secao in secoes:
            chave = secao.get("chave")
            titulo_secao = secao.get("titulo") or chave or ""
            if not chave:
                continue
            document.add_heading(_limpar(titulo_secao), level=1)
            conteudo = mapa.get(chave, extras.get(chave, ""))
            if chave == "assinaturas" and not _limpar(conteudo):
                document.add_paragraph(
                    "Assinaturas/ciência conforme o processo local:"
                )
                for item in participantes or [{"nome": "", "papel": ""}]:
                    document.add_paragraph(
                        f"{item['nome'] or '____________________________'} "
                        "_________________________________________"
                    )
            else:
                _paragrafo_texto(document, conteudo)

        document.add_paragraph("")
        rodape = document.add_paragraph(
            f"Situação da ata: {ata.get_estado_display()} · Versão interna: v{ata.version}"
            + (
                f" · Validada em {timezone.localtime(ata.validada_em):%d/%m/%Y %H:%M}"
                if ata.validada_em
                else ""
            )
        )
        for run in rodape.runs:
            run.font.size = Pt(9)

        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def render_consolidado(
        ciclo: CicloReferencial,
        *,
        salas=None,
        incluir_pendentes: bool = True,
    ) -> tuple[bytes, dict]:
        """Consolidação institucional: todos os objetivos na ordem oficial.

        Entram apenas conteúdos aprovados/publicados; GTs sem aprovação são
        marcados como pendentes, sem nenhum conteúdo inventado. Dados pessoais
        (frequência, justificativas) ficam fora por construção.
        """

        if salas is None:
            salas = (
                SalaGT.raw_objects.filter(is_deleted=False, ciclo_id=ciclo.pk)
                .select_related("cliente", "ciclo", "objetivo", "gt")
                .order_by("objetivo__ordem", "objetivo__codigo")
            )
        salas = list(salas)

        document = _novo_documento(
            ciclo.cliente_id, f"Consolidacao {ciclo.nome}"
        )
        document.add_heading("Referencial — consolidação dos objetivos", level=0)
        document.add_paragraph(
            f"Município: {ciclo.cliente.nome} · Ciclo: {_limpar(ciclo.nome)} · "
            f"Emitido em {timezone.localtime():%d/%m/%Y %H:%M}"
        )
        document.add_paragraph(
            f"Objetivos incluídos: {len(salas)}. "
            "O documento reúne apenas conteúdo aprovado ou publicado pelos GTs; "
            "objetivos ainda em elaboração aparecem como pendentes."
        )

        document.add_heading("Sumário", level=1)
        sumario = document.add_paragraph()
        _campo(sumario, 'TOC \\o "1-2" \\h \\z \\u')
        nota = document.add_paragraph(
            "Abra o documento no Word e atualize o sumário (F9) para numerar as páginas."
        )
        for run in nota.runs:
            run.italic = True
            run.font.size = Pt(9)

        trabalhos = {
            trabalho.sala_id: trabalho
            for trabalho in TrabalhoGT.raw_objects.filter(
                is_deleted=False, sala_id__in=[sala.pk for sala in salas]
            )
        }

        resumo = {"total": len(salas), "consolidados": 0, "pendentes": 0}
        for sala in salas:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            trabalho = trabalhos.get(sala.pk)
            if not incluir_pendentes and (trabalho is None or not trabalho.consolidavel):
                resumo["pendentes"] += 1
                continue
            incluiu = _secao_objetivo(document, sala, trabalho, somente_consolidavel=True)
            resumo["consolidados" if incluiu else "pendentes"] += 1

        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_heading("Situação dos Grupos de Trabalho", level=1)
        _tabela(
            document,
            ["Objetivo", "Grupo de Trabalho", "Situação do conteúdo"],
            [
                [
                    f"{sala.objetivo.codigo} — {sala.objetivo.titulo}",
                    sala.gt.nome,
                    (
                        trabalhos[sala.pk].get_estado_display()
                        if sala.pk in trabalhos
                        and trabalhos[sala.pk].consolidavel
                        else TEXTO_PENDENTE
                    ),
                ]
                for sala in salas
            ],
            larguras=[7, 5, 4],
        )

        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue(), resumo

    # ------------------------------------------------------------------
    # Registro, armazenamento e autorização
    # ------------------------------------------------------------------
    @staticmethod
    def _registrar(
        *,
        cliente_id: int,
        usuario,
        alvo_tipo: str,
        alvo_id,
        nome_arquivo: str,
        conteudo: bytes,
        metadados: dict,
    ) -> ExportJob:
        storage = private_ava_storage()
        caminho = storage.save(
            f"{PASTA_DOCUMENTOS}/{cliente_id}/{nome_arquivo}", ContentFile(conteudo)
        )
        job = ExportJob.objects.create(
            cliente_id=cliente_id,
            alvo_tipo=alvo_tipo,
            alvo_id=str(alvo_id),
            formato=ExportJob.Formato.DOCX,
            status=ExportJob.Status.DONE,
            finished_at=timezone.now(),
            payload_json={
                "arquivo": caminho,
                "nome_arquivo": nome_arquivo,
                "gerado_por": getattr(usuario, "id", None),
                **metadados,
            },
        )
        AVAAuditService.registrar(
            cliente_id=cliente_id,
            usuario=usuario,
            entidade="exports.ExportJob",
            entidade_id=job.pk,
            acao="documento_gt_gerado",
            diff={"tipo": alvo_tipo, "alvo": str(alvo_id), "arquivo": nome_arquivo},
        )
        return job

    @staticmethod
    def gerar_documento_objetivo(sala: SalaGT, usuario) -> ExportJob:
        acesso_a_sala(usuario, sala).exigir(
            "pode_exportar", "Você não pode exportar documentos deste GT."
        )
        trabalho = TrabalhoGT.raw_objects.filter(is_deleted=False, sala_id=sala.pk).first()
        conteudo = DocumentoGTService.render_objetivo(sala, trabalho)
        nome = (
            f"referencial_{_slug(sala.ciclo.nome, 24)}_obj-{_slug(sala.objetivo.codigo, 10)}"
            f"_quadro-objetivo_v{trabalho.version if trabalho else 0}"
            f"_{timezone.localdate():%Y-%m-%d}.docx"
        )
        return DocumentoGTService._registrar(
            cliente_id=sala.cliente_id,
            usuario=usuario,
            alvo_tipo=ExportJob.AlvoTipo.GT_OBJETIVO,
            alvo_id=sala.pk,
            nome_arquivo=nome,
            conteudo=conteudo,
            metadados={
                "sala": sala.pk,
                "ciclo": sala.ciclo_id,
                "objetivo": sala.objetivo.codigo,
                "estado": trabalho.estado if trabalho else None,
            },
        )

    @staticmethod
    def gerar_documento_ata(reuniao: ReuniaoGT, usuario) -> ExportJob:
        acesso_a_sala(usuario, reuniao.sala).exigir(
            "pode_exportar", "Você não pode exportar documentos deste GT."
        )
        ata = AtaReuniaoGT.raw_objects.filter(is_deleted=False, reuniao_id=reuniao.pk).first()
        if ata is None:
            raise ValidationError("Crie a minuta da ata antes de exportar.")
        conteudo = DocumentoGTService.render_ata(reuniao, ata)
        nome = (
            f"referencial_{_slug(reuniao.sala.ciclo.nome, 24)}"
            f"_obj-{_slug(reuniao.sala.objetivo.codigo, 10)}_ata"
            f"_{timezone.localtime(reuniao.data_hora):%Y-%m-%d}"
            f"_v{ata.version}.docx"
        )
        return DocumentoGTService._registrar(
            cliente_id=reuniao.cliente_id,
            usuario=usuario,
            alvo_tipo=ExportJob.AlvoTipo.GT_ATA,
            alvo_id=reuniao.pk,
            nome_arquivo=nome,
            conteudo=conteudo,
            metadados={
                "sala": reuniao.sala_id,
                "reuniao": reuniao.pk,
                "ata": ata.pk,
                "estado": ata.estado,
            },
        )

    @staticmethod
    def gerar_consolidado(
        ciclo: CicloReferencial,
        usuario,
        *,
        salas_ids: list[int] | None = None,
    ) -> ExportJob:
        """Consolidação do ciclo. Seleção parcial exige escolha explícita."""

        if not tem_visao_global(usuario):
            raise PermissionDenied(
                "Apenas a administração e o acompanhamento institucional exportam a consolidação."
            )
        if getattr(usuario, "cliente_id", None) and usuario.cliente_id != ciclo.cliente_id:
            raise PermissionDenied("Este ciclo pertence a outro município.")

        salas = (
            SalaGT.raw_objects.filter(is_deleted=False, ciclo_id=ciclo.pk)
            .select_related("cliente", "ciclo", "objetivo", "gt")
            .order_by("objetivo__ordem", "objetivo__codigo")
        )
        parcial = bool(salas_ids)
        if parcial:
            salas = salas.filter(pk__in=salas_ids)

        conteudo, resumo = DocumentoGTService.render_consolidado(ciclo, salas=salas)
        sufixo = "parcial" if parcial else "completa"
        nome = (
            f"referencial_{_slug(ciclo.nome, 24)}_consolidacao-{sufixo}"
            f"_{timezone.localdate():%Y-%m-%d}.docx"
        )
        return DocumentoGTService._registrar(
            cliente_id=ciclo.cliente_id,
            usuario=usuario,
            alvo_tipo=ExportJob.AlvoTipo.GT_CONSOLIDADO,
            alvo_id=ciclo.pk,
            nome_arquivo=nome,
            conteudo=conteudo,
            metadados={"ciclo": ciclo.pk, "parcial": parcial, "resumo": resumo},
        )

    # ------------------------------------------------------------------
    # Download
    # ------------------------------------------------------------------
    TIPOS_GT = {
        ExportJob.AlvoTipo.GT_OBJETIVO,
        ExportJob.AlvoTipo.GT_ATA,
        ExportJob.AlvoTipo.GT_CONSOLIDADO,
    }

    @staticmethod
    def autorizar_download(job: ExportJob, usuario) -> None:
        """Revalida a autorização no download — não confia na URL do arquivo."""

        if job.alvo_tipo not in DocumentoGTService.TIPOS_GT:
            raise PermissionDenied("Documento indisponível.")
        if job.alvo_tipo == ExportJob.AlvoTipo.GT_CONSOLIDADO:
            if not tem_visao_global(usuario):
                raise PermissionDenied("Você não pode baixar a consolidação institucional.")
            if getattr(usuario, "cliente_id", None) and usuario.cliente_id != job.cliente_id:
                raise PermissionDenied("Documento de outro município.")
            return

        payload = job.payload_json if isinstance(job.payload_json, dict) else {}
        sala_id = payload.get("sala")
        sala = (
            SalaGT.raw_objects.filter(is_deleted=False, pk=sala_id, cliente_id=job.cliente_id)
            .select_related("cliente", "ciclo", "objetivo", "gt")
            .first()
        )
        if sala is None:
            raise PermissionDenied("Documento indisponível.")
        acesso_a_sala(usuario, sala).exigir(
            "pode_exportar", "Você não pode baixar documentos deste GT."
        )

    @staticmethod
    def abrir_arquivo(job: ExportJob):
        payload = job.payload_json if isinstance(job.payload_json, dict) else {}
        caminho = payload.get("arquivo")
        if not caminho:
            raise FileNotFoundError("Arquivo não registrado para esta exportação.")
        return private_ava_storage().open(caminho, "rb")

    @staticmethod
    def nome_do_arquivo(job: ExportJob) -> str:
        payload = job.payload_json if isinstance(job.payload_json, dict) else {}
        return payload.get("nome_arquivo") or f"documento-{job.pk}.docx"

    @staticmethod
    def documentos_da_sala(sala: SalaGT):
        """Histórico de documentos gerados para a sala (objetivo e atas)."""

        reunioes = [
            str(pk)
            for pk in ReuniaoGT.raw_objects.filter(
                is_deleted=False, sala_id=sala.pk
            ).values_list("id", flat=True)
        ]
        filtro = Q(alvo_tipo=ExportJob.AlvoTipo.GT_OBJETIVO, alvo_id=str(sala.pk))
        if reunioes:
            filtro |= Q(alvo_tipo=ExportJob.AlvoTipo.GT_ATA, alvo_id__in=reunioes)
        return (
            ExportJob.raw_objects.filter(
                is_deleted=False,
                cliente_id=sala.cliente_id,
                formato=ExportJob.Formato.DOCX,
            )
            .filter(filtro)
            .order_by("-created_at")
        )


def resumo_frequencia_pessoal(sala: SalaGT, usuario) -> dict:
    """Frequência do próprio usuário na sala (consulta individual)."""

    presencas = PresencaReuniaoGT.raw_objects.filter(
        is_deleted=False, reuniao__sala_id=sala.pk, usuario_id=usuario.pk
    )
    total = presencas.count()
    presentes = presencas.filter(situacao=PresencaReuniaoGT.Situacao.PRESENTE).count()
    justificadas = presencas.filter(situacao=PresencaReuniaoGT.Situacao.JUSTIFICADA).count()
    return {
        "total": total,
        "presentes": presentes,
        "justificadas": justificadas,
        "ausentes": total - presentes - justificadas,
        "percentual": int(round(100 * presentes / total)) if total else 0,
    }
